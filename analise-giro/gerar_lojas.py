# -*- coding: utf-8 -*-
"""
Venda e estoque LOJA A LOJA, do banco do cliente.

A AG mostra o cliente inteiro. Esta parte abre por loja: quanto cada porta
vendeu de cada produto no mês e quanto ainda tem em estoque. É o que a aba
"AG LOJA A LOJA" da planilha faz à mão, e daqui sai pronto.

De onde vem cada coisa:
  · o universo de produtos e os rótulos (marca, cor, gênero, grupo, PDV) saem
    da mesma planilha que alimenta a AG;
  · a venda e o estoque saem do banco do cliente (Firebird), cruzando pela
    coluna COD_CLIENTE da planilha, que é a referência do produto lá dentro.

O banco é uma CÓPIA. Ele não se atualiza sozinho: quando o cliente mandar um
arquivo novo, é só rodar este script de novo.

Uso:  python gerar_lojas.py
"""
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

try:
    import fdb
except ImportError:                                   # pragma: no cover
    raise SystemExit('Falta o pacote fdb (pip install fdb).')

# o console do Windows abre em cp1252 e engasga com a seta do resumo
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

AQUI = Path(__file__).resolve().parent
PASTA = AQUI / 'dados'
ANO = 2026

BANCO = r'C:\empr01\empr01.fdb'
FBCLIENT = r'C:\Users\Administrativo\fb25\fbclient.dll'

PLANOS = Path(r'C:\Users\Administrativo\OneDrive\CMS - COMPART\OneDrive\CMS COMPART'
              r'\A0 PLANO DE CONTAS')

CLIENTES = [
    {
        'id': 'degraus',
        'nome': 'Degraus',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG DEGRAUS.xlsx',
        'aba': 'AG',
        'aba_lojas': 'AG LOJA A LOJA',     # de onde vêm os nomes das portas
    },
    {
        'id': 'degraus-ua',
        'nome': 'Degraus',
        'marca': 'Under Armour',
        'arquivo': PLANOS / 'UNDER ARMOUR' / 'JORGE_MARCOS UA' / 'AG UA DEGRAUS.xlsx',
        'aba': 'AG',
        'aba_lojas': None,                 # a de UA não tem a aba; usa a da Oly
    },
]

# Operações do ERP que contam como venda de balcão, e as que a desfazem.
# Venda no PDV entra sem código de operação (NFC-e ou sem documento fiscal).
DEVOLVE = (1, 11)        # 1 = TROCA, 11 = DEVOLUCAO DO CLIENTE
COMPRA = 3               # entrada vinda do fornecedor
LOJA_MAX = 15            # acima disso é escritório, atacado (FPS) e DPA
DESDE = 2022             # de quando em diante procurar a última entrada

CABECALHO = ('MARCA/SUB', 'MARCAS')
ALIAS = {
    'codigo': ('CODIGO',),
    'descricao': ('DESCRICAO PRODUTO', 'DESCRICAO'),
    'cor': ('COR',),
    'genero': ('GENERO',),
    'grupo': ('GRUPO_COLECAO', 'GRUPO COLECAO'),
    'pdv': ('PDV',),
    'sku': ('SKU',),
    'cod_cliente': ('COD_CLIENTE', 'REF. CLIENTE', 'REF CLIENTE'),
    'desc_cliente': ('DESC_CLIENTE', 'REF/DESC'),
}

# A referência do produto no sistema do cliente muda de plano para plano:
# na Olympikus ela vem pronta na coluna COD_CLIENTE; na Under Armour essa
# coluna costuma vir vazia e a referência está no fim da REF/DESC, que
# concatena descrição e referência ("TENIS UNDER ARMOUR SKYLINE 5
# 6014736-BKBKCR"). Pegar a cauda resolve os dois.
CAUDA_REF = re.compile(r'(\d{5,}\s*-\s*[A-Z0-9/_.]+)\s*$')


def ref_do_cliente(cod_cliente, desc_cliente):
    cc = norm(cod_cliente)
    if cc:
        return cc
    m = CAUDA_REF.search(norm(desc_cliente))
    return re.sub(r'\s+', '', m.group(1)) if m else ''



def norm(t):
    t = unicodedata.normalize('NFD', str(t or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', t).strip().upper()


def texto(v):
    return '' if v in (None, '') else str(v).strip()


def num(v):
    if v in (None, ''):
        return 0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return int(round(f)) if abs(f - round(f)) < 1e-9 else round(f, 2)


# ── planilha ──────────────────────────────────────────────────────────
def ler_produtos(ws):
    """Os produtos da AG, com o código que o cliente usa no sistema dele."""
    linhas = ws.iter_rows(values_only=True)
    cab = next((r for r in linhas if r and norm(r[0]) in CABECALHO), None)
    if cab is None:
        raise SystemExit('Cabeçalho da aba AG não encontrado.')
    col = {norm(v): i for i, v in enumerate(cab) if isinstance(v, str)}

    def onde(campo, obrigatorio=True):
        for nome in ALIAS[campo]:
            if nome in col:
                return col[nome]
        if obrigatorio:
            raise SystemExit(f'Coluna {ALIAS[campo][0]} não encontrada na aba AG.')
        return None

    ic = {c: onde(c, c not in ('grupo', 'sku', 'cod_cliente', 'desc_cliente'))
          for c in ALIAS}
    fora = []
    for r in linhas:
        if not r or not r[ic['codigo']]:
            continue
        codigo = r[ic['codigo']]
        sku = r[ic['sku']] if ic['sku'] is not None else None
        try:
            sku = int(sku)
        except (TypeError, ValueError):
            sku = None
        fora.append({
            'k': texto(r[1]) or f'{codigo}{texto(r[ic["cor"]])}',
            'marca': texto(r[0]),
            'codigo': (str(int(codigo)) if isinstance(codigo, float) and codigo.is_integer()
                       else texto(codigo)),
            'descricao': texto(r[ic['descricao']]),
            'cor': texto(r[ic['cor']]),
            'genero': texto(r[ic['genero']]).upper(),
            'grupo': texto(r[ic['grupo']]) if ic['grupo'] is not None else '',
            'pdv': num(r[ic['pdv']]) or None,
            'sku': sku,
            'ref': ref_do_cliente(
                r[ic['cod_cliente']] if ic['cod_cliente'] is not None else '',
                r[ic['desc_cliente']] if ic['desc_cliente'] is not None else ''),
        })
    return fora


def ler_carteira(wb):
    """Pares com pedido em aberto, por produto: { 'CODIGO|COR': pares }.

    É a "programação": o que já está comprado e a caminho. Produto com
    programação não entra na sugestão de agrupar — a reposição resolve.
    """
    if 'CARTEIRA' not in wb.sheetnames:
        return {}
    linhas = wb['CARTEIRA'].iter_rows(values_only=True)
    cab = next(linhas)
    col = {norm(v): i for i, v in enumerate(cab) if isinstance(v, str)}
    faltam = [c for c in ('REFERENCIA', 'COR', 'PREV FAT', 'CART') if c not in col]
    if faltam:
        return {}
    ir, ic, ip, iq = col['REFERENCIA'], col['COR'], col['PREV FAT'], col['CART']
    fora = {}
    for r in linhas:
        if not r or not r[ir] or not isinstance(r[ip], dt.datetime) or r[ip].year != ANO:
            continue
        pares = num(r[iq])
        if not pares:
            continue
        chave = f'{norm(r[ir])}|{norm(r[ic])}'
        fora[chave] = fora.get(chave, 0) + pares
    return fora


def ler_nomes_das_lojas(caminho, aba):
    """As portas têm nome na planilha ("LJ05 SHOP ILHA"); no banco são LOJA 05.

    O nome sai daqui como "LJ05 - SHOP ILHA", e esta lista é também o
    filtro: loja que não está na planilha não entra no relatório. São as
    que já não operam (03, 09, 13) e a 08, fechada desde 2025.
    """
    if not aba:
        return {}
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if aba not in wb.sheetnames:
        wb.close()
        return {}
    nomes = {}
    for r in wb[aba].iter_rows(min_row=7, max_col=7, values_only=True):
        try:
            loja = int(r[5])
        except (TypeError, ValueError):
            continue
        nome = texto(r[6])
        if nome and loja not in nomes:
            # "LJ05 SHOP ILHA" -> "LJ05 - SHOP ILHA"
            nomes[loja] = re.sub(r'^(LJ\s*\d+)\s+', lambda m: m.group(1) + ' - ', nome)
    wb.close()
    return nomes


# ── banco do cliente ──────────────────────────────────────────────────
def abrir_banco():
    return fdb.connect(dsn=BANCO, user='SYSDBA', password='masterkey',
                       fb_library_name=FBCLIENT, charset='WIN1252')


def puxar(con, sql, args=()):
    cur = con.cursor()
    cur.execute(sql, args)
    return cur.fetchall()


SQL_MOV = """
    SELECT it.FD_PRODUTO, m.FD_LOJA, EXTRACT(MONTH FROM m.FD_DATA_MOV),
           SUM(i.FD_QUANTIDADE)
      FROM TB_MOVIMENTOS m
      JOIN TB_MOVIMENTOS_ITENS i
        ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
       AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
      JOIN TB_ITENS it ON it.FD_ITEM = i.FD_ITEM
     WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
       AND {condicao}
     GROUP BY 1, 2, 3"""


def ler_banco():
    con = abrir_banco()
    inicio = dt.date(ANO, 1, 1)
    venda = defaultdict(float)      # (produto, loja, mes) -> pares
    for linha in puxar(con, SQL_MOV.format(
            condicao="m.FD_CODOPER IS NULL AND m.FD_ENTRADA_SAIDA = 'S'"), (inicio, LOJA_MAX)):
        pid, loja, mes, qt = linha
        venda[(pid, loja, int(mes))] += float(qt or 0)
    # troca e devolução do cliente voltam para a prateleira: saem da venda
    for linha in puxar(con, SQL_MOV.format(
            condicao=f"m.FD_CODOPER IN {DEVOLVE} AND m.FD_ENTRADA_SAIDA = 'E'"),
            (inicio, LOJA_MAX)):
        pid, loja, mes, qt = linha
        venda[(pid, loja, int(mes))] -= float(qt or 0)

    # Movimento LÍQUIDO do mês (tudo que entrou menos tudo que saiu): é com
    # ele que se reconstrói o estoque de cada mês fechado, andando de trás
    # para frente a partir do saldo de hoje. O banco não guarda foto mensal.
    delta = defaultdict(float)
    for pid, loja, mes, qt in puxar(con, """
        SELECT it.FD_PRODUTO, m.FD_LOJA, EXTRACT(MONTH FROM m.FD_DATA_MOV),
               SUM(CASE WHEN m.FD_ENTRADA_SAIDA = 'E' THEN i.FD_QUANTIDADE
                        ELSE -i.FD_QUANTIDADE END)
          FROM TB_MOVIMENTOS m
          JOIN TB_MOVIMENTOS_ITENS i
            ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
           AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
          JOIN TB_ITENS it ON it.FD_ITEM = i.FD_ITEM
         WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
         GROUP BY 1, 2, 3""", (inicio, LOJA_MAX)):
        delta[(pid, loja, int(mes))] += float(qt or 0)

    # Compra do ano e o dia em que o produto entrou na loja pela última vez.
    # "Última compra" aqui é a última ENTRADA: compra do fornecedor ou
    # transferência recebida — as duas põem mercadoria na prateleira.
    compra = defaultdict(float)
    for pid, loja, qt in puxar(con, """
        SELECT it.FD_PRODUTO, m.FD_LOJA, SUM(i.FD_QUANTIDADE)
          FROM TB_MOVIMENTOS m
          JOIN TB_MOVIMENTOS_ITENS i
            ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
           AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
          JOIN TB_ITENS it ON it.FD_ITEM = i.FD_ITEM
         WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
           AND m.FD_CODOPER = ? AND m.FD_ENTRADA_SAIDA = 'E'
         GROUP BY 1, 2""", (inicio, LOJA_MAX, COMPRA)):
        compra[(pid, loja)] += float(qt or 0)

    # Dia da última VENDA do produto naquela loja. É com ele que se sabe se a
    # peça está parada há mais de 30 dias — o mês fechado não serve, porque
    # no dia 24 um mês sem venda ainda pode ser só 24 dias.
    venda_ult = {}
    for pid, loja, dia in puxar(con, """
        SELECT it.FD_PRODUTO, m.FD_LOJA, MAX(m.FD_DATA_MOV)
          FROM TB_MOVIMENTOS m
          JOIN TB_MOVIMENTOS_ITENS i
            ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
           AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
          JOIN TB_ITENS it ON it.FD_ITEM = i.FD_ITEM
         WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
           AND m.FD_CODOPER IS NULL AND m.FD_ENTRADA_SAIDA = 'S'
         GROUP BY 1, 2""", (dt.date(DESDE, 1, 1), LOJA_MAX)):
        venda_ult[(pid, loja)] = dia

    ultima = {}
    for pid, loja, dia in puxar(con, """
        SELECT it.FD_PRODUTO, m.FD_LOJA, MAX(m.FD_DATA_MOV)
          FROM TB_MOVIMENTOS m
          JOIN TB_MOVIMENTOS_ITENS i
            ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
           AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
          JOIN TB_ITENS it ON it.FD_ITEM = i.FD_ITEM
         WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
           AND m.FD_ENTRADA_SAIDA = 'E' AND (m.FD_CODOPER IS NULL OR m.FD_CODOPER <> 11)
         GROUP BY 1, 2""", (dt.date(DESDE, 1, 1), LOJA_MAX)):
        ultima[(pid, loja)] = dia

    # Venda da loja inteira, todas as marcas: é ela que diz se a porta estava
    # aberta no mês. Sem isso, loja que não vendeu Under Armour em setembro
    # pareceria fechada — e a UA vende pouco, então seria quase toda.
    loja_mes = defaultdict(float)
    for loja, mes, qt in puxar(con, """
        SELECT m.FD_LOJA, EXTRACT(MONTH FROM m.FD_DATA_MOV), SUM(i.FD_QUANTIDADE)
          FROM TB_MOVIMENTOS m
          JOIN TB_MOVIMENTOS_ITENS i
            ON i.FD_MOVIMENTO_ID = m.FD_MOVIMENTO_ID
           AND i.FD_MOVIMENTO_DG = m.FD_MOVIMENTO_DG
         WHERE m.FD_DATA_MOV >= ? AND m.FD_LOJA <= ?
           AND m.FD_CODOPER IS NULL AND m.FD_ENTRADA_SAIDA = 'S'
         GROUP BY 1, 2""", (inicio, LOJA_MAX)):
        loja_mes[(loja, int(mes))] += float(qt or 0)

    estoque = defaultdict(float)
    for pid, loja, saldo in puxar(con, 'SELECT FD_PRODUTO, FD_LOJA, FD_SALDO '
                                       'FROM TB_SALDOS_PRODUTOS '
                                       'WHERE FD_SALDO <> 0 AND FD_LOJA <= ?', (LOJA_MAX,)):
        estoque[(pid, loja)] += float(saldo or 0)

    # referência do cliente -> produto, e o dia do último movimento (idade da cópia)
    por_ref, por_id = {}, set()
    for pid, ref in puxar(con, 'SELECT FD_PRODUTO, TRIM(FD_REFERENCIA) FROM TB_PRODUTOS'):
        por_id.add(pid)
        if ref:
            por_ref.setdefault(ref.upper(), pid)
    ate = puxar(con, 'SELECT MAX(FD_DATA_MOV) FROM TB_MOVIMENTOS')[0][0]
    con.close()
    return {'venda': venda, 'estoque': estoque, 'delta': delta, 'compra': compra,
            'ultima': ultima, 'venda_ult': venda_ult, 'loja_mes': loja_mes,
            'por_ref': por_ref, 'por_id': por_id, 'ate': ate}


# ── junta tudo ────────────────────────────────────────────────────────
def saldo_por_mes(estoque_hoje, deltas, n):
    """Estoque no fim de cada mês, de trás para frente a partir de hoje.

    saldo(m-1) = saldo(m) - (o que entrou menos o que saiu no mês m).
    """
    # O estoque negativo do ERP (venda lançada sem a entrada correspondente)
    # fica como está: é assim que a planilha do cliente mostra, e esconder
    # isso faria a conta fechar errado no mês seguinte.
    saldos = [0] * n
    atual = estoque_hoje
    for i in range(n - 1, -1, -1):
        saldos[i] = int(round(atual))
        atual -= deltas[i]
    return saldos


def gerar(cli, banco, nomes_lojas):
    venda, estoque = banco['venda'], banco['estoque']
    delta, compra, ultima = banco['delta'], banco['compra'], banco['ultima']
    loja_mes, venda_ult = banco['loja_mes'], banco['venda_ult']
    por_ref, por_id, ate = banco['por_ref'], banco['por_id'], banco['ate']
    print(f'· {cli["nome"]} {cli["marca"]}: lendo {cli["arquivo"].name}')
    wb = openpyxl.load_workbook(cli['arquivo'], read_only=True, data_only=True)
    produtos = ler_produtos(wb[cli['aba']])
    carteira = ler_carteira(wb)
    wb.close()

    meses = [f'{ANO}-{m:02d}' for m in range(1, ate.month + 1)]
    # quem não entrou e o motivo: a AG usa isso para avisar antes do clique
    #   sem-codigo   → a planilha não tem a referência do cliente (COD_CLIENTE)
    #   sem-cadastro → tem a referência, mas ela não existe no banco dele
    #   sem-giro     → está cadastrado e nunca teve venda, estoque ou compra
    saida, faltando, lojas_vistas, usados = [], {}, set(), set()
    for p in produtos:
        pid = p['sku'] if p['sku'] in por_id else por_ref.get(p['ref'])
        # Último recurso, só para linha sem cor (vestuário, que o cliente
        # cadastra pelo artigo): o próprio código. Com cor isso juntaria
        # cores diferentes no mesmo produto.
        if not pid and not p['cor']:
            candidato = por_ref.get(norm(p['codigo']))
            if candidato and candidato not in usados:
                pid = candidato
        if not pid:
            faltando[p['k']] = 'sem-codigo' if (not p['ref'] and not p['sku']) else 'sem-cadastro'
            continue
        usados.add(pid)
        lojas = {}
        for loja in sorted(nomes_lojas):
            n = ate.month
            v = [int(round(venda.get((pid, loja, m), 0))) for m in range(1, n + 1)]
            e = int(round(estoque.get((pid, loja), 0)))
            cp = int(round(compra.get((pid, loja), 0)))
            uc = ultima.get((pid, loja))
            uv = venda_ult.get((pid, loja))
            if not any(v) and not e and not cp:
                continue
            ds = [delta.get((pid, loja, m), 0) for m in range(1, n + 1)]
            reg = {'s': saldo_por_mes(e, ds, n)}
            if any(v):
                reg['v'] = v
            if e:
                reg['e'] = e
            if cp:
                reg['c'] = cp
            if uc:
                reg['uc'] = uc.isoformat()
            if uv:
                reg['uv'] = uv.isoformat()
            lojas[str(loja)] = reg
            lojas_vistas.add(loja)
        if not lojas:
            faltando[p['k']] = 'sem-giro'
            continue
        prog = carteira.get(f"{norm(p['codigo'])}|{norm(p['cor'])}", 0)
        saida.append({**{c: p[c] for c in ('k', 'marca', 'codigo', 'descricao', 'cor',
                                           'genero', 'grupo', 'pdv')},
                      'sku': pid, 'ref': p['ref'], **({'prog': prog} if prog else {}),
                      'l': lojas})

    base = {
        'cliente': {'id': cli['id'], 'nome': cli['nome'], 'marca': cli['marca']},
        'ano': ANO,
        'meses': meses,
        'atualizadoEm': dt.datetime.now().isoformat(timespec='seconds'),
        'banco': {'arquivo': Path(BANCO).name, 'ate': ate.isoformat()},
        # 'v' é a venda da LOJA INTEIRA no mês (todas as marcas): serve para
        # saber se a porta estava aberta, e para medir o peso da marca lá
        'lojas': [{'id': l, 'nome': nomes_lojas[l],
                   'v': [int(round(loja_mes.get((l, m), 0))) for m in range(1, ate.month + 1)]}
                  for l in sorted(lojas_vistas)],
        'produtos': saida,
        'faltando': faltando,
    }
    PASTA.mkdir(exist_ok=True)
    destino = PASTA / f'lojas-{cli["id"]}.json'
    destino.write_text(json.dumps(base, ensure_ascii=False, separators=(',', ':')),
                       encoding='utf-8')
    pares = sum(sum(r.get('v', [])) for p in saida for r in p['l'].values())
    est = sum(r.get('e', 0) for p in saida for r in p['l'].values())
    comprog = sum(1 for x in saida if x.get('prog'))
    print(f'  {comprog} produtos com programação em aberto')
    motivos = {}
    for m in faltando.values():
        motivos[m] = motivos.get(m, 0) + 1
    print(f'  {len(saida)} produtos em {len(lojas_vistas)} lojas · '
          f'{pares:,} pares vendidos · {est:,} em estoque'.replace(',', '.'))
    print('  fora: ' + ', '.join(f'{n} {m}' for m, n in sorted(motivos.items())))
    print(f'  → {destino.name} ({destino.stat().st_size // 1024} KB)')


def main():
    print(f'Banco: {BANCO}')
    banco = ler_banco()
    print(f'  movimento até {banco["ate"]:%d/%m/%Y}')
    nomes = {}
    for cli in CLIENTES:
        nomes.update(ler_nomes_das_lojas(cli['arquivo'], cli['aba_lojas']))
    for cli in CLIENTES:
        gerar(cli, banco, nomes)


if __name__ == '__main__':
    main()
