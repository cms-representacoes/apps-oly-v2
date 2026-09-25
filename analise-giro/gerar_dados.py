# -*- coding: utf-8 -*-
"""
Análise de Giro — gera o banco de dados (JSON) a partir do plano de contas.

Cada cliente tem o seu plano (.xlsx) com a aba AG: uma linha por produto e,
para cada mês, um bloco de seis colunas — Venda · Estoque · Giro · Cob. ·
Análise · OBS. Daqui sai só o ano corrente (2026), e só o que a tela mostra
nesta primeira etapa: venda, estoque e OBS por mês. Giro e cobertura não são
guardados: a tela calcula a partir de venda e estoque, então nunca ficam
desatualizados.

A carteira que a tela mostra não sai daqui: ela é lida da Detalhada, que se
atualiza todo dia, pelos `codigos` do cliente. A aba CARTEIRA do plano ainda
entra no JSON como registro.

Uso:  python gerar_dados.py            (todos os clientes)
      python gerar_dados.py degraus    (só um)
"""
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

import shutil
import tempfile

# o console do Windows abre em cp1252 e engasga com a seta do resumo
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ANO = 2026
PASTA = Path(__file__).resolve().parent / 'dados'
# as fotos dos produtos, as mesmas da Carteira e da Detalhada
FOTOS = Path(__file__).resolve().parent.parent / 'imagens_olympikus'
PLANOS = Path(r'C:\Users\Administrativo\OneDrive\CMS - COMPART\OneDrive\CMS COMPART\A0 PLANO DE CONTAS')

# Um por cliente. O COD_CLIENTE é a chave que casa o relatório do cliente
# com a linha do plano (é por ele que hoje se faz o PROCV), então ele vai
# para o banco mesmo sem ser usado ainda.
CLIENTES = [
    {
        'id': 'degraus',
        'nome': 'Degraus',
        'razao': 'DEGRAUS FPS',
        'codigo': '3834566',
        # códigos do cliente na Detalhada: em 2026 a Degraus trocou a filial de
        # recebimento, e a carteira dela vem pelos dois
        'codigos': ['3834566', '124330'],
        'lider': '104988',
        'lojas': 11,
        'vendedor': 'JORGE',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG DEGRAUS.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
    {
        'id': 'degraus-ua',
        'nome': 'Degraus',
        'razao': 'DEGRAUS FPS',
        'codigo': '3834566',
        'codigos': ['3834566', '124330'],
        'lider': '104988',
        'lojas': 9,
        'vendedor': 'JORGE',
        'marca': 'Under Armour',
        'arquivo': PLANOS / 'UNDER ARMOUR' / 'JORGE_MARCOS UA' / 'AG UA DEGRAUS.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
    {
        'id': 'talentus',
        'nome': 'Talentus',
        'razao': 'TALENTUS',
        # A planilha abre com "CLIENTE: 17759", que é um código antigo: na
        # Detalhada a Talentus compra pelos códigos abaixo, os dois com o
        # JORGE. É por eles que a carteira vem.
        'codigo': '18529',
        'codigos': ['18529', '65218'],
        'lider': '3616654',
        'lojas': 10,
        'vendedor': 'JORGE',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG TALENTUS.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
    {
        'id': 'fit-store',
        'nome': 'Fit Store',
        'razao': 'FIT STORE BY COMBATE',
        'codigo': '2245855',
        'codigos': ['2245855'],
        'lider': '',
        'lojas': 1,
        'vendedor': 'JORGE',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG FITSTORE.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
    {
        'id': 'ideal-magazine',
        'nome': 'Ideal Magazine',
        'razao': 'GRUPO IDEAL',
        # o grupo compra por 15 códigos na Detalhada; o 21511 é o líder
        'codigo': '21511',
        'codigos': ['21511', '11948', '105709', '68621', '121069', '25931', '63095',
                    '75383', '76687', '95206', '137388', '1026378', '2373565',
                    '3132954', '3309898'],
        'lider': '21511',
        'lojas': 7,
        'vendedor': 'JORGE',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG IDEAL MAGAZINE.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
    {
        'id': 'matriz-esportes',
        'nome': 'Matriz Esportes',
        'razao': 'MATRIZ ESPORTE',
        'codigo': '60437',
        'codigos': ['60437', '1948498', '3068604'],
        'lider': '',
        'lojas': 3,
        'vendedor': 'JORGE',
        'marca': 'Olympikus',
        'arquivo': PLANOS / 'OLYMPIKUS' / 'JORGE' / 'AG MATRIZ ESPORTE.xlsx',
        'aba': 'AG',
        'aba_carteira': 'CARTEIRA',
    },
]


def abrir(caminho):
    """Abre a planilha mesmo com ela aberta no Excel (lê uma cópia)."""
    try:
        return openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except PermissionError:
        copia = Path(tempfile.gettempdir()) / f'ag_{abs(hash(str(caminho)))}.xlsx'
        shutil.copy2(caminho, copia)
        print('  (planilha aberta no Excel: li uma cópia)')
        return openpyxl.load_workbook(copia, read_only=True, data_only=True)


def norm(t):
    t = unicodedata.normalize('NFD', str(t or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', t).strip().upper()


def num(v):
    if v in (None, ''):
        return 0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return int(round(f)) if abs(f - round(f)) < 1e-9 else round(f, 2)


def texto(v):
    return '' if v in (None, '') else str(v).strip()


# O plano da Olympikus abre a linha do cabeçalho com MARCA/SUB e o da
# Under Armour com MARCAS; o código do cliente é COD_CLIENTE num e
# REF. CLIENTE no outro. Fora isso os dois têm a mesma forma.
CABECALHO = ('MARCA/SUB', 'MARCAS')
# A mesma família aparece com nomes diferentes de plano para plano: a Matriz
# escreve só FTW, a Fit Store escreve OLY MEIA, a Talentus OLY CHIN. Todas
# viram o nome que os chips da tela conhecem.
SUB_ALIAS = {
    'FTW': 'OLY FTW', 'VEST': 'OLY VEST', 'MEIA': 'OLY MEIAS', 'MEIAS': 'OLY MEIAS',
    'ACC': 'OLY ACC', 'CHI': 'OLY CHI', 'CHIN': 'OLY CHI',
    'OLY MEIA': 'OLY MEIAS', 'OLY CHIN': 'OLY CHI',
}


def sub_marca(valor):
    v = norm(valor)
    return SUB_ALIAS.get(v, v)
ALIAS = {
    'codigo': ('CODIGO',),
    'descricao': ('DESCRICAO PRODUTO', 'DESCRICAO'),
    'cor': ('COR',),
    'genero': ('GENERO',),
    'grupo': ('GRUPO_COLECAO', 'GRUPO COLECAO'),
    'pdv': ('PDV',),
    'cod_cliente': ('COD_CLIENTE', 'REF. CLIENTE', 'REF CLIENTE', 'COD CLIENTE'),
    'desc_cliente': ('DESC_CLIENTE', 'REF/DESC', 'DESC. CLIENTE'),
    # a coluna que junta código e cor numa chave só; na Degraus ela é a
    # segunda coluna, sem título, e na Talentus se chama UPLOAD
    'chave': ('UPLOAD', 'CONCATENAR', 'CONCA', 'CHAVE'),
}

# A referência do produto no sistema do cliente muda de plano para plano:
# na Olympikus ela vem pronta na coluna COD_CLIENTE; na Under Armour essa
# coluna costuma vir vazia e a referência está no fim da REF/DESC, que
# concatena descrição e referência ("TENIS UNDER ARMOUR SKYLINE 5
# 6014736-BKBKCR"). Pegar a cauda resolve os dois.
CAUDA_REF = re.compile(r'(\d{5,}\s*-\s*[A-Z0-9/_.]+)\s*$')


def ref_do_cliente(cod_cliente, desc_cliente):
    cc = norm(cod_cliente)
    if cc:
        return cc
    m = CAUDA_REF.search(norm(desc_cliente))
    return re.sub(r'\s+', '', m.group(1)) if m else ''



def ler_ag(ws, marca='Olympikus'):
    """Produtos e meses do ano, lidos pelo cabeçalho da aba.

    `marca` diz qual plano é: linha de outra marca na planilha fica de fora.
    A Ideal e a Matriz têm um punhado de linhas UA no meio do plano Olympikus.
    """
    # O cabeçalho nem sempre começa na coluna A: na Ideal e na Matriz a
    # primeira coluna é a CONCA, e o MARCA/SUB vem depois.
    linhas = ws.iter_rows(values_only=True)
    cab = None
    for r in linhas:
        if r and any(norm(v) in CABECALHO for v in r if isinstance(v, str)):
            cab = r
            break
    if cab is None:
        raise SystemExit('Cabeçalho da aba AG não encontrado '
                         f'(esperava {" ou ".join(CABECALHO)} em alguma coluna).')

    col = {norm(v): i for i, v in enumerate(cab) if isinstance(v, str)}

    def onde(campo, obrigatorio=True):
        for nome in ALIAS[campo]:
            if nome in col:
                return col[nome]
        if obrigatorio:
            raise SystemExit(f'Coluna {ALIAS[campo][0]} não encontrada na aba AG.')
        return None

    ic = {
        'marca': next(col[n] for n in CABECALHO if n in col),
        'chave': onde('chave', False) if 'chave' in ALIAS else 1,
        'codigo': onde('codigo'),
        'descricao': onde('descricao'),
        'cor': onde('cor'),
        'genero': onde('genero'),
        'grupo': onde('grupo', False),
        'pdv': onde('pdv'),
        'cod_cliente': onde('cod_cliente', False),
        'desc_cliente': onde('desc_cliente', False),
    }
    if ic['chave'] is None:
        ic['chave'] = 1
    # blocos de mês: a célula do cabeçalho é a data e é também a coluna de
    # venda; as seguintes são estoque, giro, cob., análise e OBS
    meses = []
    for i, v in enumerate(cab):
        if isinstance(v, dt.datetime) and v.year == ANO and i + 5 < len(cab) \
                and norm(cab[i + 1]) == 'ESTOQUE' and norm(cab[i + 5]) == 'OBS':
            meses.append((f'{v.year}-{v.month:02d}', i))
    if not meses:
        raise SystemExit(f'Nenhum mês de {ANO} encontrado na aba AG.')

    ehUA = marca.upper().startswith('UNDER')
    produtos = []
    fora_da_marca = 0
    for r in linhas:
        if not r or not r[ic['codigo']]:
            continue
        sub = sub_marca(r[ic['marca']])
        if sub.startswith('UA') != ehUA:
            fora_da_marca += 1
            continue
        m = {}
        for chave, i in meses:
            venda, estoque, obs = num(r[i]), num(r[i + 1]), texto(r[i + 5])
            reg = {}
            if venda:
                reg['v'] = venda
            if estoque:
                reg['e'] = estoque
            if obs:
                reg['o'] = obs
            if reg:
                m[chave] = reg
        # só entra quem teve venda ou estoque no ano; OBS sozinha não conta
        if not any(x.get('v') or x.get('e') for x in m.values()):
            continue
        codigo = r[ic['codigo']]
        produtos.append({
            'k': texto(r[ic['chave']]) or f'{codigo}{texto(r[ic["cor"]])}',
            'marca': sub,
            'codigo': str(int(codigo)) if isinstance(codigo, float) and codigo.is_integer() else texto(codigo),
            'descricao': texto(r[ic['descricao']]),
            'cor': texto(r[ic['cor']]),
            'genero': texto(r[ic['genero']]).upper(),
            'grupo': texto(r[ic['grupo']]) if ic['grupo'] is not None else '',
            'pdv': num(r[ic['pdv']]) or None,
            'codCliente': ref_do_cliente(
                r[ic['cod_cliente']] if ic['cod_cliente'] is not None else '',
                r[ic['desc_cliente']] if ic['desc_cliente'] is not None else ''),
            'm': m,
        })
    if fora_da_marca:
        print(f'  {fora_da_marca} linhas de outra marca ficaram de fora')
    return [c for c, _ in meses], produtos


def ler_carteira(ws):
    """Previsão de faturamento de 2026: { 'REF|COR': { 'AAAA-MM': pares } }."""
    linhas = ws.iter_rows(values_only=True)
    cab = next(linhas)
    col = {norm(v): i for i, v in enumerate(cab) if isinstance(v, str)}
    ir, ic, ip, iq = col['REFERENCIA'], col['COR'], col['PREV FAT'], col['CART']
    cart = {}
    for r in linhas:
        if not r or not r[ir] or not isinstance(r[ip], dt.datetime) or r[ip].year != ANO:
            continue
        pares = num(r[iq])
        if not pares:
            continue
        chave = f'{texto(r[ir])}|{texto(r[ic])}'
        mes = f'{r[ip].year}-{r[ip].month:02d}'
        cart.setdefault(chave, {})
        cart[chave][mes] = cart[chave].get(mes, 0) + pares
    return cart


def gerar(cli):
    print(f'· {cli["nome"]}: lendo {cli["arquivo"].name}')
    wb = abrir(cli['arquivo'])
    meses, produtos = ler_ag(wb[cli['aba']], cli['marca'])
    carteira = ler_carteira(wb[cli['aba_carteira']]) if cli.get('aba_carteira') in wb.sheetnames else {}
    wb.close()

    agora = dt.datetime.now().isoformat(timespec='seconds')
    base = {
        'cliente': {k: cli[k] for k in ('id', 'nome', 'razao', 'codigo', 'codigos', 'lider', 'lojas', 'vendedor', 'marca')},
        'ano': ANO,
        'meses': meses,
        'atualizadoEm': agora,
        'fonte': cli['arquivo'].name,
        'produtos': produtos,
        'carteira': carteira,
    }
    PASTA.mkdir(exist_ok=True)
    destino = PASTA / f'{cli["id"]}.json'
    destino.write_text(json.dumps(base, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'  {len(produtos)} produtos · meses {meses[0]} a {meses[-1]} · '
          f'{len(carteira)} itens de carteira → {destino.name} ({destino.stat().st_size // 1024} KB)')
    return {
        'id': cli['id'], 'nome': cli['nome'], 'codigo': cli['codigo'], 'codigos': cli['codigos'], 'lojas': cli['lojas'],
        'vendedor': cli['vendedor'], 'marca': cli['marca'], 'ano': ANO, 'meses': meses,
        'produtos': len(produtos), 'atualizadoEm': agora,
    }


def indice_fotos():
    """Índice das fotos que existem: { 'ARTIGO|COR': 'arquivo.jpg' }.

    A cor entra sem separador nenhum porque os arquivos não seguem um padrão:
    a mesma cor aparece como PTO_LM, PTOLM e PTO/LM. Variações do mesmo
    produto (__02, -3000x3000.raw) ficam de fora — a tela quer uma foto por
    cor, e a principal é a que não tem sufixo.
    """
    if not FOTOS.is_dir():
        print('! pasta de fotos não encontrada:', FOTOS)
        return 0
    idx = {}
    for arq in sorted(FOTOS.iterdir()):
        if arq.suffix.lower() not in ('.jpg', '.jpeg', '.png', '.webp'):
            continue
        nome = arq.stem
        if '__' in nome or '-3000x3000' in nome.lower():
            continue                      # variação da mesma foto
        artigo, _, cor = nome.partition('_')
        if not artigo or not cor:
            continue
        chave = f'{artigo.upper()}|{re.sub(r"[^A-Z0-9]", "", cor.upper())}'
        idx.setdefault(chave, arq.name)   # a primeira em ordem alfabética manda
    destino = PASTA / 'fotos.json'
    PASTA.mkdir(exist_ok=True)
    destino.write_text(json.dumps(idx, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'fotos: {len(idx)} cores com imagem → {destino.name} ({destino.stat().st_size // 1024} KB)')
    return len(idx)


def main():
    so = set(a.lower() for a in sys.argv[1:])
    indice_arq = PASTA / 'clientes.json'
    indice = {}
    if indice_arq.exists():
        indice = {c['id']: c for c in json.loads(indice_arq.read_text(encoding='utf-8'))}
    for cli in CLIENTES:
        if so and cli['id'] not in so:
            continue
        indice[cli['id']] = gerar(cli)
    indice_fotos()
    lista = sorted(indice.values(), key=lambda c: c['nome'])
    indice_arq.write_text(json.dumps(lista, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'índice: {len(lista)} cliente(s) em {indice_arq.name}')


if __name__ == '__main__':
    main()
